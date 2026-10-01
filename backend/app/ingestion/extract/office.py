"""Word (via Pandoc) and PowerPoint (via python-pptx) to Markdown."""

import re
import subprocess
from pathlib import Path

import pypandoc
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.base import BaseShape
from pptx.slide import Slide

from app.core.config import IngestionConfig
from app.ingestion.extract.common import ExtractedPage, VisionReason, markdown_table

PANDOC_TIMEOUT_SECONDS = 120
# Pandoc writes embedded images as links into a media folder we do not keep.
_IMAGE_LINK = re.compile(r"!\[([^\]]*)\]\([^)]*\)(\{[^}]*\})?")


def extract_docx(path: Path, config: IngestionConfig) -> list[ExtractedPage]:
    """Word equations (OMML) come out as $...$ LaTeX, which is why Pandoc is
    used rather than a plain-text reader."""
    result = subprocess.run(  # noqa: S603  # fixed binary and arguments, no shell
        [
            pypandoc.get_pandoc_path(),
            str(path),
            "--from=docx",
            # Plain $...$ / $$...$$ maths, which the renderer expects (gfm would
            # write GitHub-style $`...`$ and ```math blocks instead).
            "--to=commonmark+tex_math_dollars+pipe_tables",
            "--wrap=none",
        ],
        capture_output=True,
        timeout=PANDOC_TIMEOUT_SECONDS,
        check=True,
    )
    markdown = _IMAGE_LINK.sub(
        lambda m: f"[figure{': ' + m.group(1) if m.group(1) else ''}]",
        result.stdout.decode("utf-8"),
    )
    return [ExtractedPage(1, markdown.strip(), 0.0, None)]


def _shape_markdown(shape: BaseShape) -> list[str]:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [line for child in shape.shapes for line in _shape_markdown(child)]  # type: ignore[attr-defined]
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = [[cell.text for cell in row.cells] for row in shape.table.rows]  # type: ignore[attr-defined]
        return [markdown_table(rows[0], rows[1:])] if rows else []
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return ["[figure]"]
    if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
        return []
    lines = []
    for paragraph in shape.text_frame.paragraphs:  # type: ignore[attr-defined]
        text = "".join(run.text for run in paragraph.runs).strip()
        if text:
            lines.append("  " * paragraph.level + "- " + text)
    return lines


def _has_equations(slide: Slide) -> bool:
    """Office maths (OMML) is invisible to python-pptx; detect it in the XML."""
    xml = slide._element.xml  # python-pptx exposes no public accessor
    return "oMath" in xml or "<a14:m>" in xml


def extract_pptx(path: Path, config: IngestionConfig) -> list[ExtractedPage]:
    pages: list[ExtractedPage] = []
    for index, slide in enumerate(Presentation(str(path)).slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None else ""
        body: list[str] = [f"## {title}"] if title else []
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            body.extend(_shape_markdown(shape))
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            notes = notes_frame.text.strip() if notes_frame is not None else ""
            if notes:
                body.append(f"**Speaker notes:** {notes}")
        reason = VisionReason.EQUATION_OBJECTS if _has_equations(slide) else None
        pages.append(ExtractedPage(index, "\n\n".join(body), 1.0 if reason else 0.0, reason))
    return pages
