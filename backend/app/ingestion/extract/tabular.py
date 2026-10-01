"""Spreadsheets and CSV: a schema, a row sample and numeric summaries per
sheet — not every row, which would flood search with numbers."""

import csv
import math
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from app.core.config import IngestionConfig
from app.ingestion.extract.common import ExtractedPage, markdown_table

MAX_COLUMNS_SHOWN = 30


@dataclass
class _ColumnStats:
    count: int = 0
    total: float = 0.0
    minimum: float = math.inf
    maximum: float = -math.inf
    non_numeric: int = 0

    def add(self, value: object) -> None:
        if value is None or value == "":
            return
        try:
            number = float(value)  # type: ignore[arg-type]
        except (TypeError, ValueError):
            self.non_numeric += 1
            return
        if math.isnan(number):
            return
        self.count += 1
        self.total += number
        self.minimum = min(self.minimum, number)
        self.maximum = max(self.maximum, number)


def _fmt(number: float) -> str:
    return f"{number:.6g}"


def summarise(name: str, rows: Iterable[Sequence[object]], sample_rows: int) -> str:
    iterator = iter(rows)
    header: list[object] = []
    for row in iterator:
        if any(v not in (None, "") for v in row):
            header = [v if v not in (None, "") else f"column {i + 1}" for i, v in enumerate(row)]
            break
    if not header:
        return f"## {name}\n\n(empty)"

    header = header[:MAX_COLUMNS_SHOWN]
    stats = [_ColumnStats() for _ in header]
    sample: list[Sequence[object]] = []
    count = 0
    for row in iterator:
        if not any(v not in (None, "") for v in row):
            continue
        count += 1
        if len(sample) < sample_rows:
            sample.append(row[: len(header)])
        for column, value in zip(stats, row, strict=False):
            column.add(value)

    numeric = [
        (header[i], s) for i, s in enumerate(stats) if s.count and s.non_numeric <= s.count * 0.05
    ]
    parts = [
        f"## {name}",
        f"{count} rows × {len(header)} columns. Columns: {', '.join(str(h) for h in header)}.",
        f"First {len(sample)} rows:\n\n{markdown_table(header, sample)}" if sample else "",
    ]
    if numeric:
        summary = [
            [h, s.count, _fmt(s.minimum), _fmt(s.maximum), _fmt(s.total / s.count)]
            for h, s in numeric
        ]
        parts.append(
            "Numeric columns:\n\n"
            + markdown_table(["column", "values", "min", "max", "mean"], summary)
        )
    return "\n\n".join(p for p in parts if p)


def extract_xlsx(path: Path, config: IngestionConfig) -> list[ExtractedPage]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return [
            ExtractedPage(
                index,
                summarise(
                    f"Sheet: {sheet.title}",
                    sheet.iter_rows(values_only=True),
                    config.sheet_sample_rows,
                ),
                0.0,
                None,
            )
            for index, sheet in enumerate(workbook.worksheets, start=1)
        ]
    finally:
        workbook.close()


def extract_csv(path: Path, config: IngestionConfig) -> list[ExtractedPage]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        sample = handle.read(64 * 1024)
        handle.seek(0)
        try:
            dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
                sample, delimiters=",;\t|"
            )
        except csv.Error:
            dialect = csv.excel
        rows = csv.reader(handle, dialect)
        return [ExtractedPage(1, summarise("Data", rows, config.sheet_sample_rows), 0.0, None)]
