"""Build real test files in code, so no copyrighted material is committed.

The "damaged maths" page reproduces what the design doc observed in real
lecture notes: equations set in a maths font whose extracted text is debris.
"""

import io
import subprocess
import zipfile
from pathlib import Path

import pymupdf
import pypandoc
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

PROSE = (
    "Integration by parts follows from the product rule for derivatives. "
    "It is useful when the integrand is a product of two functions, one of which "
    "becomes simpler when differentiated and the other easy to integrate."
)
# The component-wise product from the design doc, as plain extraction left it.
DEBRIS = ["1 × 2", "2 3 4", "× × ×", "−1", "= =", "2 1", "+ −", "3 × 4"]


def pdf(path: Path, pages: list[str] | None = None, *, heading: str | None = None) -> Path:
    doc = pymupdf.open()
    for text in pages or [PROSE]:
        page = doc.new_page()
        y = 72
        if heading:
            page.insert_text((72, y), heading, fontsize=22)
            y += 40
        page.insert_textbox(pymupdf.Rect(72, y, 520, 780), text, fontsize=11)
    doc.save(path)
    doc.close()
    return path


def damaged_maths_pdf(path: Path) -> Path:
    """Page 1 is prose; page 2 is an equation laid out in a maths font."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_textbox(pymupdf.Rect(72, 72, 520, 780), PROSE, fontsize=11)
    page = doc.new_page()
    page.insert_text((72, 72), "Vectors", fontsize=11)
    for i, debris in enumerate(DEBRIS):
        page.insert_text((90 + (i % 2) * 40, 110 + i * 18), debris, fontsize=12, fontname="symb")
    doc.save(path)
    doc.close()
    return path


def scanned_pdf(path: Path) -> Path:
    """A page that is only an image: no text layer."""
    image = Image.new("RGB", (600, 800), "white")
    buffer = io.BytesIO()
    image.save(buffer, "PNG")
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(page.rect, stream=buffer.getvalue())
    doc.new_page()  # and a genuinely blank page
    doc.save(path)
    doc.close()
    return path


def encrypted_pdf(path: Path) -> Path:
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "secret")
    aes256 = pymupdf.PDF_ENCRYPT_AES_256  # type: ignore[attr-defined]  # runtime constant
    doc.save(path, encryption=aes256, user_pw="pw", owner_pw="pw")
    doc.close()
    return path


def docx(path: Path, markdown: str) -> Path:
    subprocess.run(  # noqa: S603  # test helper, fixed binary
        [pypandoc.get_pandoc_path(), "--from=markdown", "--to=docx", "-o", str(path)],
        input=markdown.encode(),
        check=True,
    )
    return path


def with_extra_zip_entry(source: Path, path: Path, name: str, data: bytes) -> Path:
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            out.writestr(item, src.read(item.filename))
        out.writestr(name, data)
    return path


def pptx(path: Path) -> Path:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Eigenvalues"
    body = slide.placeholders[1].text_frame
    body.text = "Definition"
    sub = body.add_paragraph()
    sub.text = "Av = λv for non-zero v"
    sub.level = 1
    slide.notes_slide.notes_text_frame.text = "Stress the non-zero condition."
    table_slide = deck.slides.add_slide(deck.slide_layouts[5])
    table_slide.shapes.title.text = "Rates"
    table = table_slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    for (r, c), value in {(0, 0): "term", (0, 1): "rate", (1, 0): "1y", (1, 1): "4.5%"}.items():
        table.cell(r, c).text = value
    deck.save(str(path))
    return path


def xlsx(path: Path) -> Path:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Prices"
    sheet.append(["day", "price"])
    for day, price in enumerate([100.0, 101.5, 99.0, 102.25], start=1):
        sheet.append([day, price])
    book.save(path)
    return path


def jpeg_with_gps(path: Path) -> Path:
    image = Image.new("RGB", (40, 20), "red")
    exif = Image.Exif()
    exif[0x0112] = 6  # orientation: rotate 90° when displayed
    exif[0x8825] = {1: "N", 2: (53.0, 24.0, 21.0)}  # GPS: Liverpool
    image.save(path, "JPEG", exif=exif.tobytes())
    return path
